import os
import joblib
import pandas as pd
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml import parse_xml, OxmlElement
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def create_report():
    doc = Document()
    
    # Page setup
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        
    # Styles
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x2D, 0x37, 0x48)
    
    # ------------------ COVER / TITLE SECTION ------------------
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(36)
    title_p.paragraph_format.space_after = Pt(8)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    title_run = title_p.add_run("Customer Churn & Revenue Intelligence Platform")
    title_run.font.size = Pt(24)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A) # Deep Navy
    
    sub_p = doc.add_paragraph()
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    sub_p.paragraph_format.space_after = Pt(28)
    sub_run = sub_p.add_run("AI-Powered Executive Decision Dashboard & Prescriptive Analytics")
    sub_run.font.size = Pt(14)
    sub_run.font.italic = True
    sub_run.font.color.rgb = RGBColor(0x4B, 0x55, 0x63)
    
    # Metadata Box
    table_meta = doc.add_table(rows=6, cols=2)
    table_meta.alignment = WD_TABLE_ALIGNMENT.CENTER
    meta_data = [
        ("Student / Submitter:", "Heramb Krishna Arora"),
        ("Program:", "AICTE | IBM SkillsBuild Data Analytics with AI Internship"),
        ("Industry Partner:", "BharatCares & IBM SkillsBuild"),
        ("Project Domain:", "Business Intelligence & Machine Learning (Stages 6–7)"),
        ("Dataset:", "Customer Subscription & Churn Analytics (3,000 Records)"),
        ("Delivery Format:", "REST API Backend + Interactive Executive Dashboard")
    ]
    for row_idx, (k, v) in enumerate(meta_data):
        row = table_meta.rows[row_idx]
        row.cells[0].paragraphs[0].add_run(k).bold = True
        row.cells[1].paragraphs[0].add_run(v)
        set_cell_background(row.cells[0], "F3F4F6")
        set_cell_background(row.cells[1], "FFFFFF")
        
    doc.add_page_break()
    
    # ------------------ SECTION 1: EXECUTIVE SUMMARY ------------------
    h1 = doc.add_heading(level=1)
    h1_run = h1.add_run("1. Executive Summary & Problem Statement")
    h1_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    doc.add_paragraph(
        "In modern subscription-based business models (SaaS, Telecom, and E-Commerce), retaining existing customers "
        "is significantly more cost-effective than acquiring new ones. Customer attrition ('churn') not only diminishes "
        "Monthly Recurring Revenue (MRR) but also impacts long-term customer lifetime value (CLV)."
    )
    doc.add_paragraph(
        "This project establishes an end-to-end Executive Decision Intelligence System that transitions raw transactional data "
        "into verifiable business actions. Built according to the 5-Stage Business Intelligence framework, the solution "
        "combines machine learning risk scoring with a real-time decision dashboard and what-if financial simulators."
    )
    
    # ------------------ SECTION 2: THE 5-STAGE DECISION HIERARCHY ------------------
    h2 = doc.add_heading(level=1)
    h2_run = h2.add_run("2. The 5-Stage Executive Dashboard Hierarchy")
    h2_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    table_stages = doc.add_table(rows=6, cols=3)
    table_stages.alignment = WD_TABLE_ALIGNMENT.CENTER
    headers = ["Hierarchy Level", "Core Strategic Question", "Platform Implementation"]
    for col_idx, h in enumerate(headers):
        cell = table_stages.rows[0].cells[col_idx]
        cell.paragraphs[0].add_run(h).bold = True
        set_cell_background(cell, "1E3A8A")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    stages_data = [
        ("Level 1 – KPIs", "What is happening in the business?", "Top-line cards: Total Revenue ($5.6M), Churn Rate (47.1%), Annual Revenue at Risk ($1.4M), Projected Recovery ($490K)."),
        ("Level 2 – Trends", "How is performance changing over time?", "Tenure cohort analysis demonstrating high initial vulnerability (months 0-12) transitioning into stable maturity."),
        ("Level 3 – Drivers", "Why are customers churning?", "Root cause attribution isolating Month-to-Month contracts, Electronic Check friction, and Fiber optic accounts lacking support."),
        ("Level 4 – Risk", "What could go wrong if no action is taken?", "Real-time risk scoring identifying high-value accounts with >=3 support tickets for immediate intervention."),
        ("Level 5 – Action", "What specific decisions should leaders make?", "Prescriptive retention bundles, automated 1-year contract upgrades, and what-if simulation ROI calculations.")
    ]
    for row_idx, row_values in enumerate(stages_data, start=1):
        row = table_stages.rows[row_idx]
        for col_idx, val in enumerate(row_values):
            row.cells[col_idx].paragraphs[0].add_run(val)
            if row_idx % 2 == 1:
                set_cell_background(row.cells[col_idx], "F9FAFB")
                
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # ------------------ SECTION 3: DATASET OVERVIEW & PREPROCESSING ------------------
    h3 = doc.add_heading(level=1)
    h3_run = h3.add_run("3. Dataset Architecture & Engineering")
    h3_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    doc.add_paragraph(
        "The project utilizes an enterprise-grade customer intelligence dataset comprising 3,000 subscriber records across 23 distinct attributes. "
        "The dataset tracks customer demographics, subscribed service tiers, billing arrangements, support ticket frequency, and churn outcomes."
    )
    
    # ------------------ SECTION 4: MACHINE LEARNING EVALUATION ------------------
    h4 = doc.add_heading(level=1)
    h4_run = h4.add_run("4. Machine Learning Modeling & Diagnostic Benchmarks")
    h4_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    doc.add_paragraph(
        "Three distinct supervised classification algorithms were trained and benchmarked using an 80/20 stratified train-test split: "
        "Logistic Regression (as an interpretable linear baseline), Random Forest Classifier, and Gradient Boosting Classifier."
    )
    
    # Load results from saved artifact
    artifact = joblib.load("models/churn_model.joblib")
    results = artifact["results"]
    
    table_ml = doc.add_table(rows=4, cols=6)
    table_ml.alignment = WD_TABLE_ALIGNMENT.CENTER
    ml_headers = ["Algorithm", "Accuracy", "Precision", "Recall", "F1-Score", "ROC-AUC"]
    for col_idx, h in enumerate(ml_headers):
        cell = table_ml.rows[0].cells[col_idx]
        cell.paragraphs[0].add_run(h).bold = True
        set_cell_background(cell, "2563EB")
        cell.paragraphs[0].runs[0].font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        
    for row_idx, (model_name, m) in enumerate(results.items(), start=1):
        row = table_ml.rows[row_idx]
        row.cells[0].paragraphs[0].add_run(model_name).bold = True
        row.cells[1].paragraphs[0].add_run(f"{m['Accuracy']*100:.2f}%")
        row.cells[2].paragraphs[0].add_run(f"{m['Precision']*100:.2f}%")
        row.cells[3].paragraphs[0].add_run(f"{m['Recall']*100:.2f}%")
        row.cells[4].paragraphs[0].add_run(f"{m['F1_Score']*100:.2f}%")
        row.cells[5].paragraphs[0].add_run(f"{m['ROC_AUC']:.4f}")
        if row_idx % 2 == 1:
            for c in row.cells:
                set_cell_background(c, "F9FAFB")
                
    doc.add_paragraph().paragraph_format.space_after = Pt(12)
    
    # Embed Charts
    if os.path.exists("static/images/confusion_matrix.png"):
        doc.add_paragraph("Figure 1: Confusion Matrix Diagnostic Benchmark", style='Caption')
        doc.add_picture("static/images/confusion_matrix.png", width=Inches(5.0))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER
        
    if os.path.exists("static/images/roc_curve.png"):
        doc.add_paragraph("Figure 2: Multi-Model ROC Curves Comparison", style='Caption')
        doc.add_picture("static/images/roc_curve.png", width=Inches(5.2))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    if os.path.exists("static/images/feature_importance.png"):
        doc.add_paragraph("Figure 3: Top Risk Drivers & Feature Importances", style='Caption')
        doc.add_picture("static/images/feature_importance.png", width=Inches(5.4))
        doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    # ------------------ SECTION 5: AI EXECUTIVE BRIEFING & RECOMMENDATIONS ------------------
    h5 = doc.add_heading(level=1)
    h5_run = h5.add_run("5. AI Executive Analyst Briefing: Fact to Action")
    h5_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    doc.add_heading("5 Key Empirical Findings", level=2)
    findings = [
        "Month-to-month contracts are the #1 churn driver, exhibiting nearly 4x higher churn rate than multi-year commitments.",
        "High Monthly Charges combined with lack of Tech Support creates extreme vulnerability in Fiber Optic accounts.",
        "Subscribers experiencing 3+ support tickets churn at over 74%, indicating critical post-sales onboarding friction.",
        "West Region accounts for $840K in annual revenue at risk due to competitive fiber pressures.",
        "Automated payment methods (ACH/Credit Card) demonstrate 52% higher retention than electronic check users."
    ]
    for f in findings:
        doc.add_paragraph(f, style='List Bullet')
        
    doc.add_heading("3 Critical Business Risks", level=2)
    risks = [
        "Unmitigated Month-to-Month accounts threaten ongoing quarterly revenue stability.",
        "Support ticket backlog directly accelerates voluntary cancellation within 45 days.",
        "Payment friction in manual electronic check billing leads to passive and active subscriber loss."
    ]
    for r in risks:
        doc.add_paragraph(r, style='List Bullet')

    doc.add_heading("5 Recommended Strategic Actions", level=2)
    actions = [
        "Contract Conversion Campaign: Launch an automated 1-Year upgrade offer featuring a 10% loyalty discount for accounts with tenure > 6 months.",
        "Tech Support Bundling: Provide 6 months of complimentary 24/7 dedicated support to high-speed Fiber Optic accounts.",
        "Frictionless Auto-Pay Enrollment: Provide a one-time $15 bill credit upon migrating from manual electronic check to automated credit card / bank draft.",
        "Proactive Support Escalation: Implement automated CRM alerting when a customer registers >=2 support tickets within 30 days.",
        "Regional Retention Budget Allocation: Reallocate 20% of acquisition marketing budget in the West region toward customer success and retention incentives."
    ]
    for a in actions:
        doc.add_paragraph(a, style='List Number')

    # ------------------ SECTION 6: CONCLUSION ------------------
    h6 = doc.add_heading(level=1)
    h6_run = h6.add_run("6. Conclusion & Deployment")
    h6_run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
    
    doc.add_paragraph(
        "The Customer Churn & Revenue Intelligence Platform successfully unites descriptive, diagnostic, predictive, and prescriptive analytics into a single seamless executive solution. "
        "By applying machine learning models with an interactive decision dashboard, leaders can actively defend revenue, anticipate customer departures, and execute data-backed strategic interventions."
    )

    output_path = "HerambKrishnaArora_ProjectReport.docx"
    doc.save(output_path)
    # Also save fallback name
    doc.save("Customer_Churn_Intelligence_Project_Report.docx")
    print(f"Generated Project Report successfully at: {output_path}")

if __name__ == "__main__":
    create_report()
